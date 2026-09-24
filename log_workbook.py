"""Builds the downloadable .xlsx run log (per-event log, run summary, raw console log)."""

from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from config import LOG_COLUMN_HEADERS


def write_log_workbook(path_or_stream, event_log, log_lines, stats):
    """path_or_stream: a filesystem path OR a file-like object (e.g. io.BytesIO)."""
    wb = Workbook()
    log_ws = wb.active
    log_ws.title = "Log"
    headers = [
        LOG_COLUMN_HEADERS["name"],
        LOG_COLUMN_HEADERS["id"],
        LOG_COLUMN_HEADERS["date"],
        LOG_COLUMN_HEADERS["event"],
        LOG_COLUMN_HEADERS["reason"],
        LOG_COLUMN_HEADERS["details"],
    ]
    log_ws.append(headers)

    records = sorted(
        event_log.values(),
        key=lambda record: (
            record["date"] is None,
            record["date"] or date.max,
            record["id"],
            record["name"],
        ),
    )
    for record in records:
        log_ws.append(
            [
                record["name"],
                record["id"],
                record["date"],
                ", ".join(record["events"]),
                "\n".join(record["reasons"]),
                "\n".join(record["details"]),
            ]
        )

    summary_ws = wb.create_sheet("Run summary")
    summary_ws.append(["Metric", "Value"])
    for metric, value in (
        ("Cells marked", stats["marked"]),
        ("Weekend N cells", stats["weekend_n"]),
        ("Overlaps combined", stats["overlaps_combined"]),
        ("Already marked (skipped)", stats["already_marked"]),
        ("Conflicts noted", stats["conflicts"]),
        ("IDs not in roster", stats["no_id_match"]),
        ("No date overlap", stats["no_date_overlap"]),
        ("Name/ID mismatches", stats["name_mismatch"]),
    ):
        summary_ws.append([metric, value])

    console_ws = wb.create_sheet("Console log")
    console_ws.append(["Message"])
    for line in log_lines:
        console_ws.append([line])

    header_fill = PatternFill("solid", fgColor="1F4E78")
    for ws in (log_ws, summary_ws, console_ws):
        for cell in ws[1]:
            cell.font = Font(color="FFFFFFFF", bold=True)
            cell.fill = header_fill
        ws.freeze_panes = "A2"
        ws.sheet_view.showGridLines = False

    for row in log_ws.iter_rows(min_row=2):
        row[2].number_format = "dd/mm/yyyy"
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    log_ws.auto_filter.ref = log_ws.dimensions
    log_ws.column_dimensions["A"].width = 28
    log_ws.column_dimensions["B"].width = 15
    log_ws.column_dimensions["C"].width = 13
    log_ws.column_dimensions["D"].width = 20
    log_ws.column_dimensions["E"].width = 60
    log_ws.column_dimensions["F"].width = 42

    summary_ws.column_dimensions["A"].width = 25
    summary_ws.column_dimensions["B"].width = 18
    console_ws.column_dimensions["A"].width = 120
    for row in console_ws.iter_rows(min_row=2):
        row[0].alignment = Alignment(wrap_text=True, vertical="top")

    wb.save(path_or_stream)
